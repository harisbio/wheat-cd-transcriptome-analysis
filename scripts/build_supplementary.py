
import csv, glob, gzip, json, os, shutil, sys
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Cm, Pt

Image.MAX_IMAGE_PIXELS = None
ROOT, WR, SRA, V2, UNIPROT = sys.argv[1:6]
OUT = os.path.join(V2, "Supplementary_Material")
FIG, TAB, CODE = (os.path.join(OUT, d) for d in ("Supplementary_Figures", "Supplementary_Tables", "Data_and_Code"))
for d in (FIG, TAB, CODE):
    os.makedirs(d, exist_ok=True)
J = os.path.join


def rcsv(p):
    with open(p, encoding="utf-8") as f:
        return list(csv.DictReader(f))


ann = {r["gene"]: r for r in csv.DictReader(open(J(ROOT, "wheat_gene_annotation.tsv"), encoding="utf-8"), delimiter="\t")}

# ======================================================================= figures
legends_fig = {
 "S1": "Wheat seedling growth under the ten treatment combinations (T1–T10). Rows show successive growth stages (top to bottom) and columns the individual treatments; labels on the pots indicate biological replicates (R1–R3).",
 "S2": "Treatment of textile effluent with chemically synthesized (FeO(C)-NPs) and biogenic (FeO(B)-NPs) iron oxide nanoparticles. (a) Effluent after addition of FeO(C)-NPs or FeO(B)-NPs (1 mg mL⁻¹) and 7 h incubation under sunlight, with the untreated control. (b) The same samples after centrifugation (10,000 rpm, 10 min, 4 °C) to remove nanoparticles. Decolorization was greatest with FeO(B)-NPs.",
 "S3": "Chemical oxygen demand (COD) determination in textile effluent treated with FeO(B)-NPs and FeO(C)-NPs. (a) Reaction mixtures before digestion: distilled-water control, untreated effluent, effluent + FeO(B)-NPs and effluent + FeO(C)-NPs. (b) The same mixtures after digestion; the color change of the nanoparticle-treated samples reflects their lower COD.",
 "S4": "Phosphate determination in textile effluent treated with FeO(B)-NPs and FeO(C)-NPs. Blue color intensity is proportional to phosphate concentration: untreated effluent (WW) > effluent treated with FeO(C)-NPs (CNP WW) > effluent treated with FeO(B)-NPs (BNP WW).",
 "S5": "Sulfate determination in textile effluent treated with FeO(B)-NPs and FeO(C)-NPs. (a) Distilled-water control, (b) untreated effluent, (c) effluent treated with FeO(C)-NPs and (d) effluent treated with FeO(B)-NPs. Turbidity is proportional to residual sulfate.",
 "S6": "Wheat seedlings (T1–T10) used for length measurements. (a) Shoot length and (b) root length measurement for all ten treatments, with a ruler for scale.",
 "S7": "Transcriptomic response of wheat shoots to Cd stress (reanalysis of BioProject PRJNA1124581; wild-type shoots, n = 3 per condition). (A) Density of log₂-transformed normalized counts. (B) Principal component analysis of variance-stabilized expression (500 most variable genes). (C) Numbers of differentially expressed genes (DEGs; adjusted p < 0.05, |log₂FC| ≥ 1.5). (D) Volcano plot of shrunken log₂FC against −log₁₀ adjusted p-value. (E) MA plot. (F) Circular heatmap of the 1,000 most variable genes (row-scaled variance-stabilized expression); rings are samples, the outer band indicates condition.",
 "S8": "GO enrichment of (A) downregulated and (B) upregulated DEGs in Cd-treated wheat shoots. Ten non-redundant terms per ontology (BP, CC, MF) are shown; dot size indicates the number of DEGs and color the −log₁₀ FDR (hypergeometric test, Benjamini–Hochberg FDR < 0.05).",
 "S9": "KEGG pathway enrichment of shoot DEGs. (A) Significantly enriched pathways (FDR < 0.05) ranked by −log₁₀ FDR; labels give the number of DEGs with up/down counts; global and overview maps are not shown. (B) Chord diagram linking the seven most significant pathways to their eight most significant DEGs, colored by log₂FC.",
 "S10": "Functional gene groups among shoot DEGs (oxidative stress, transcription factors and transporters; GO-based definitions as in Fig. 7). (A–C) Circular heatmaps of the 30 most significant DEGs per group. (D–F) Violin plots of the expression of all DEGs in each group per sample.",
 "S11": "Protein–protein interaction networks of shoot DEGs (STRING v12.5, combined score ≥ 0.4). (A–C) Circular networks of the oxidative-stress, transcription factor and transporter DEGs grouped into family sectors; node color = log₂FC, node size = degree. (D) The 20 top hub genes of the shoot stress-response network ranked by Maximal Clique Centrality. (E) Circular heatmap of hub-gene expression.",
}
src_fig = {f"S{i}": J(V2, n) for i, n in enumerate(["S1.png", "S2.png", "S3.jpeg", "S4.png", "S5.png", "S6.png"], 1)}
src_fig.update({
 "S7": J(ROOT, "FigS_Shoot_supplementary", "Composite_Shoot.tif"),
 "S8": J(ROOT, "FigS_GO_Shoot_supplementary", "Composite_GO_Shoot.tif"),
 "S9": J(ROOT, "FigS_KEGG_Shoot_supplementary", "Composite_KEGG_Shoot.tif"),
 "S10": J(ROOT, "FigS_GeneGroups_Shoot_supplementary", "Composite_GeneGroups_Shoot.tif"),
 "S11": J(ROOT, "FigS_PPI_Shoot_supplementary", "Composite_PPI_Shoot.tif"),
})
previews = {}
for k, src in src_fig.items():
    dst = J(FIG, f"Figure {k}.tif")
    im = Image.open(src).convert("RGB")
    if src.lower().endswith((".tif", ".tiff")):
        shutil.copy2(src, dst)
    else:
        dpi = im.info.get("dpi", (300, 300))
        im.save(dst, compression="tiff_lzw", dpi=(max(300, round(dpi[0])),) * 2)
    pv = J(OUT, f"_preview_{k}.jpg")
    p = im.copy(); p.thumbnail((2400, 2400)); p.save(pv, quality=88)
    previews[k] = pv
print("figures:", len(src_fig))

# ======================================================================= tables
HFILL = PatternFill("solid", fgColor="DDE7F0")


def write_sheet(ws, title, note, header, rows):
    ws.append([title]); ws["A1"].font = Font(bold=True, size=12)
    ws.append([note]); ws["A2"].alignment = Alignment(wrap_text=False)
    ws.append([])
    ws.append(header)
    for c in ws[4]:
        c.font = Font(bold=True); c.fill = HFILL
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A5"
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(9, width + 2))


def book(name, sheets):
    wb = Workbook(); wb.remove(wb.active)
    for (sname, title, note, header, rows) in sheets:
        write_sheet(wb.create_sheet(sname[:31]), title, note, header, rows)
    wb.save(J(TAB, name))


def num(x, nd=4):
    try:
        v = float(x)
        return round(v, nd) if abs(v) >= 1e-3 or v == 0 else float(f"{v:.3e}")
    except (TypeError, ValueError):
        return x


# ---- Table S3: libraries and QC
key = {"SRR29482939": ("Root", "Control", 3, "CK-WT-R3"), "SRR29482940": ("Root", "Control", 2, "CK-WT-R2"),
       "SRR29482941": ("Root", "Control", 1, "CK-WT-R1"), "SRR29482948": ("Root", "Cd", 3, "T-WT-R3"),
       "SRR29482950": ("Root", "Cd", 2, "T-WT-R2"), "SRR29482951": ("Root", "Cd", 1, "T-WT-R1"),
       "SRR29482938": ("Shoot", "Control", 1, "CK-WT-S1"), "SRR29482959": ("Shoot", "Control", 2, "CK-WT-S2"),
       "SRR29482958": ("Shoot", "Control", 3, "CK-WT-S3"), "SRR29482947": ("Shoot", "Cd", 1, "T-WT-S1"),
       "SRR29482946": ("Shoot", "Cd", 2, "T-WT-S2"), "SRR29482945": ("Shoot", "Cd", 3, "T-WT-S3")}
sal = {r["Sample"].replace("trimmed_", "").replace("_quant", ""): r for r in csv.DictReader(open(J(SRA, "salmon_QC_summary.txt")), delimiter="\t")}
s3 = []
for run, (tis, cond, rep, title) in sorted(key.items(), key=lambda t: (t[1][0], t[1][1] != "Control", t[1][2])):
    fp = json.load(open(J(SRA, "fastp_results", f"{run}_fastp.json")))["summary"]
    b, a = fp["before_filtering"], fp["after_filtering"]
    s3.append([run, title, "PRJNA1124581", tis, "Wild type", cond, rep, b["total_reads"], a["total_reads"],
               round(100 * a["total_reads"] / b["total_reads"], 2), round(100 * a["q30_rate"], 2),
               int(sal[run]["Total_Fragments"]), int(sal[run]["Mapped_Fragments"]), float(sal[run]["Mapping_Rate(%)"])])
book("Table S3.xlsx", [("Table S3", "Table S3. RNA-seq libraries reanalyzed in this study and read-processing statistics.",
      "Wild-type wheat roots and shoots from BioProject PRJNA1124581 (Cheng et al., 2025). Reads: fastp v0.23.4 (both mates counted); fragments: Salmon v1.10.2 selective alignment to IWGSC RefSeq v1.1 cDNA (Ensembl Plants release 61).",
      ["SRA run", "Sample title (SRA)", "BioProject", "Tissue", "Genotype", "Condition", "Replicate", "Raw reads",
       "Clean reads", "Reads retained (%)", "Q30 bases, clean (%)", "Salmon fragments", "Mapped fragments", "Mapping rate (%)"], s3)])

# ---- Table S4/S5: DEGs
groups = {}
for tis, gf in (("Root", "Fig7_GeneGroups_Root_main"), ("Shoot", "FigS_GeneGroups_Shoot_supplementary")):
    for g in ("Oxidative_stress", "Transcription_factors", "Transporters"):
        for r in rcsv(J(ROOT, gf, f"{tis}_{g}_DEGs.csv")):
            groups.setdefault((tis, r["gene"]), []).append(g.replace("_", " "))


def deg_rows(tis):
    out = []
    for r in rcsv(J(ROOT, f"{tis}_DEGs.csv")):
        g = r[""]; a = ann.get(g, {})
        out.append([g, a.get("uniprot", ""), a.get("protein_name", ""), num(r["baseMean"], 2), num(r["log2FC"], 3),
                    num(r["log2FC_unshrunk"], 3), num(r["lfcSE"], 3), num(r["pvalue"]), num(r["padj"]), r["sig"],
                    "; ".join(groups.get((tis, g), []))])
    return out


hdr = ["Gene ID (IWGSC v1.1)", "UniProt accession", "Protein name (UniProt)", "Base mean", "log2FC (apeglm-shrunken)",
       "log2FC (unshrunk)", "lfcSE", "p-value", "Adjusted p-value (BH)", "Direction", "Functional group (Fig. 7 / S10)"]
for tn, tis in (("S4", "Root"), ("S5", "Shoot")):
    rows = deg_rows(tis)
    up = sum(r[9] == "Upregulated" for r in rows)
    book(f"Table {tn}.xlsx", [(f"Table {tn}", f"Table {tn}. Differentially expressed genes in Cd-treated vs control wild-type wheat {tis.lower()}s ({len(rows)} DEGs; {up} up, {len(rows) - up} down).",
          "DESeq2 with apeglm log2FC shrinkage; DEGs: adjusted p < 0.05 and |shrunken log2FC| >= 1.5. Protein names from UniProtKB (taxon 4565). Sorted by adjusted p-value.", hdr, rows)])

# ---- Table S6: GO
go_sheets = []
for tis, gf in (("Root", "Fig5_GO_Root_main"), ("Shoot", "FigS_GO_Shoot_supplementary")):
    for d in ("Up", "Down"):
        nr = {r["GO"] + r["Ontology"] for r in rcsv(J(ROOT, gf, f"{tis}_GO_{d}_nonredundant.csv"))}
        rows = [[r["Ontology"], r["GO"], r["Term"], int(r["Count"]), int(r["SetSize"]), num(r["GeneRatio"]),
                 num(r["FoldEnrichment"], 3), num(r["pvalue"]), num(r["padj"]), "yes" if r["GO"] + r["Ontology"] in nr else "no",
                 r["Genes"]] for r in rcsv(J(ROOT, gf, f"{tis}_GO_{d}_all_significant.csv"))]
        go_sheets.append((f"{tis}_{d}", f"Table S6 ({tis}, {d.lower()}regulated). Significantly enriched GO terms (FDR < 0.05).",
                          "One-sided hypergeometric test against all tested genes with GO annotation; terms with 10-500 genes; Benjamini-Hochberg FDR. 'Non-redundant' = retained after Jaccard (>=0.7) filtering and eligible for plotting.",
                          ["Ontology", "GO ID", "Term", "DEGs in term", "Background genes in term", "Gene ratio",
                           "Fold enrichment", "p-value", "FDR", "Non-redundant", "DEG IDs"], rows))
book("Table S6.xlsx", go_sheets)

# ---- Table S7: KEGG
kg = []
for tis, kf in (("Root", "Fig6_KEGG_Root_main"), ("Shoot", "FigS_KEGG_Shoot_supplementary")):
    up = {r[""] for r in rcsv(J(ROOT, f"{tis}_DEGs.csv")) if float(r["log2FC"]) > 0}
    for d in ("All", "Up", "Down"):
        rows = []
        for r in rcsv(J(ROOT, kf, f"{tis}_KEGG_{d}_significant.csv")):
            g = r["Genes"].split(";") if r["Genes"] else []
            nu = sum(x in up for x in g)
            rows.append([r["Pathway"].replace("taes", "map"), r["Name"], "yes" if r["Overview"] == "TRUE" else "no",
                         int(r["Count"]), nu, len(g) - nu, int(r["SetSize"]), num(r["FoldEnrichment"], 3),
                         num(r["pvalue"]), num(r["padj"]), r["Genes"]])
        kg.append((f"{tis}_{d}DEGs", f"Table S7 ({tis}, {d.lower()} DEGs). Significantly enriched KEGG pathways (FDR < 0.05).",
                   "KEGG REST (organism taes, accessed 13 Oct 2026) via UniProt mapping; one-sided hypergeometric test; background = tested genes mapped to >= 1 pathway; pathways with >= 5 genes; BH FDR.",
                   ["KEGG map", "Pathway", "Global/overview map", "DEGs in pathway", "Up", "Down", "Background genes",
                    "Fold enrichment", "p-value", "FDR", "DEG IDs"], rows))
book("Table S7.xlsx", kg)

# ---- Table S8: functional groups
g8 = []
for tis, gf in (("Root", "Fig7_GeneGroups_Root_main"), ("Shoot", "FigS_GeneGroups_Shoot_supplementary")):
    summ = rcsv(J(ROOT, gf, f"{tis}_gene_group_summary.csv"))
    for g in ("Oxidative_stress", "Transcription_factors", "Transporters"):
        rows = [[r["gene"], ann.get(r["gene"], {}).get("uniprot", ""), r["ProteinName"], num(r["baseMean"], 2),
                 num(r["log2FC"], 3), num(r["padj"]), r["Direction"]] for r in rcsv(J(ROOT, gf, f"{tis}_{g}_DEGs.csv"))]
        s = [x for x in summ if x["Group"] == g.replace("_", " ")][0]
        g8.append((f"{tis}_{g}"[:31], f"Table S8 ({tis}, {g.replace('_', ' ')}). {s['DEGs']} DEGs: {s['Up']} up, {s['Down']} down ({s['Percent_up']}% up).",
                   "Groups defined by GO annotation including descendant terms: oxidative stress GO:0016209, GO:0006979, GO:0004364; transcription factors GO:0003700; transporters GO:0005215.",
                   ["Gene ID", "UniProt accession", "Protein name", "Base mean", "log2FC (shrunken)", "Adjusted p-value", "Direction"], rows))
book("Table S8.xlsx", g8)

# ---- Table S9: PPI hubs (stress network) + network nodes
g9 = []
for tis, pf in (("Root", "Fig8_PPI_Root_main"), ("Shoot", "FigS_PPI_Shoot_supplementary")):
    rows = [[int(r["rank"]), r["gene"], r["string_id"], r["protein"], num(r["log10_MCC"], 3), int(float(r["degree"])),
             num(r["log2FC"], 3), num(r["padj"]), r["group"]] for r in rcsv(J(ROOT, pf, f"{tis}_hub_genes_MCC_stress_network.csv"))]
    g9.append((f"{tis}_MCC_ranking", f"Table S9 ({tis}). MCC ranking of all proteins in the stress-response PPI network (top 20 shown in Fig. {'8D' if tis == 'Root' else 'S11D'}).",
               "STRING v12.5 (taxon 4565), interactions with combined score >= 0.4 among the union of oxidative-stress, transcription factor and transporter DEGs. MCC as in cytoHubba (sum over maximal cliques of (|C|-1)!), reported as log10.",
               ["MCC rank", "Gene ID", "STRING ID", "Protein name", "log10(MCC)", "Degree", "log2FC", "Adjusted p-value", "Functional group"], rows))
    for g in ("Oxidative_stress", "Transcription_factors", "Transporters"):
        rows = [[r["gene"], r["protein"], int(float(r["degree"])), num(r["log2FC"], 3)] for r in rcsv(J(ROOT, pf, f"{tis}_PPI_{g}_nodes.csv"))]
        g9.append((f"{tis}_{g}_nodes"[:31], f"Table S9 ({tis}, {g.replace('_', ' ')} network). Network nodes (Fig. {'8' if tis == 'Root' else 'S11'}{'ABC'[['Oxidative_stress','Transcription_factors','Transporters'].index(g)]}).",
                   "Degree within the group network (STRING combined score >= 0.4). Proteins with degree 0 were not drawn.",
                   ["Gene ID", "Protein (abbreviated)", "Degree", "log2FC"], rows))
book("Table S9.xlsx", g9)

# ---- Table S10: all-DEG high-confidence MCC
g10 = []
for tis, pf in (("Root", "Fig8_PPI_Root_main"), ("Shoot", "FigS_PPI_Shoot_supplementary")):
    rows = [[int(r["rank"]), r["gene"], r["string_id"], r["protein"], num(r["log10_MCC"], 3), int(float(r["degree"])),
             num(r["log2FC"], 3), num(r["padj"])] for r in rcsv(J(ROOT, pf, f"{tis}_hub_genes_MCC_allDEGs_score0.7_top200.csv"))]
    g10.append((f"{tis}_allDEG_MCC_top200", f"Table S10 ({tis}). Top 200 proteins by MCC in the high-confidence network of all {tis.lower()} DEGs.",
                "STRING v12.5, combined score >= 0.7, all DEGs. Shown for comparison with the stress-response network; top-ranked hubs are dominated by ribosomal proteins and ribosome-assembly factors.",
                ["MCC rank", "Gene ID", "STRING ID", "Protein name", "log10(MCC)", "Degree", "log2FC", "Adjusted p-value"], rows))
book("Table S10.xlsx", g10)

# ---- Table S11: Fig. 9 family values
rows = []
for r in rcsv(J(ROOT, "Fig9_Model_Root", "fig9_family_values.csv")):
    n, up, dn = int(r["n_DEGs"]), int(r["n_up"]), int(r["n_down"])
    resp = "mixed" if min(up, dn) / n >= 0.3 else ("up" if float(r["mean_log2FC"]) > 0 else "down")
    rows.append([r["family"], r["regex_on_UniProt_name"], n, up, dn, num(r["mean_log2FC"], 2), resp, r["genes"]])
book("Table S11.xlsx", [("Table S11", "Table S11. Gene-family values shown in the Cd-response model of wheat root cells (Fig. 9).",
      "Root DEGs assigned to families by matching UniProt protein names (regular expression shown). Mixed = >= 30% of DEGs in each direction. Counts can differ from the GO-based groups of Fig. 7.",
      ["Family (Fig. 9 label)", "Matching rule (UniProt protein name)", "DEGs", "Up", "Down", "Mean log2FC", "Response", "Gene IDs"], rows)])

# ---- Tables S1-S2 (existing Word file)
shutil.copy2(J(V2, "Supplementry file 2.docx"), J(TAB, "Tables S1-S2.docx"))
print("tables: S1-S2 (docx), S3-S11 (xlsx)")

# ======================================================================= Supplementary_Material.docx
tab_leg = {
 "S1": ("Tables S1-S2.docx", "Comparison of FeO-NP-mediated physiological recovery in wheat with published iron oxide nanoparticle plant studies."),
 "S2": ("Tables S1-S2.docx", "Comparison of the wastewater-remediation efficiency of FeO-NPs with published iron oxide nanoparticle studies."),
 "S3": ("Table S3.xlsx", "RNA-seq libraries reanalyzed (BioProject PRJNA1124581, wild-type roots and shoots), with read-processing and mapping statistics."),
 "S4": ("Table S4.xlsx", "Differentially expressed genes in Cd-treated vs control wheat roots, with UniProt annotation and functional-group membership."),
 "S5": ("Table S5.xlsx", "Differentially expressed genes in Cd-treated vs control wheat shoots."),
 "S6": ("Table S6.xlsx", "Significantly enriched GO terms for up- and downregulated DEGs in roots and shoots, with the DEGs in each term."),
 "S7": ("Table S7.xlsx", "Significantly enriched KEGG pathways for all, up- and downregulated DEGs in roots and shoots."),
 "S8": ("Table S8.xlsx", "DEGs in the oxidative-stress, transcription factor and transporter groups (roots and shoots)."),
 "S9": ("Table S9.xlsx", "Maximal Clique Centrality ranking of the stress-response PPI networks and node lists of the group networks (roots and shoots)."),
 "S10": ("Table S10.xlsx", "Top 200 MCC-ranked proteins in the high-confidence (score ≥ 0.7) PPI networks of all root and shoot DEGs."),
 "S11": ("Table S11.xlsx", "Gene-family values used in the root-cell model (Fig. 9), with the matching rule and gene identifiers."),
}
d = docx.Document()
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
d.add_heading("Supplementary Material", 0)
d.add_paragraph("Biogenic Iron Oxide Nanoparticles Remediate Textile Wastewater and Alleviate Its Phytotoxicity in Wheat, "
                "with Insights from Public Cadmium-Stress Transcriptome Data").runs[0].italic = True
d.add_heading("Supplementary Figures", 1)
for k in [f"S{i}" for i in range(1, 12)]:
    d.add_picture(previews[k], width=Cm(16))
    p = d.add_paragraph(); r = p.add_run(f"Figure {k}. "); r.bold = True; p.add_run(legends_fig[k])
    p.paragraph_format.space_after = Pt(14)
d.add_heading("Supplementary Tables", 1)
for k, (fname, leg) in tab_leg.items():
    p = d.add_paragraph(); r = p.add_run(f"Table {k}. "); r.bold = True; p.add_run(leg + f" [{fname}]")
d.save(J(OUT, "Supplementary_Material.docx"))
for pv in previews.values():
    os.remove(pv)
print("Supplementary_Material.docx written")

# ======================================================================= Data_and_Code (Zenodo)
sub = {k: J(CODE, k) for k in ("scripts", "input", "annotation", "results", "qc")}
for p in sub.values():
    os.makedirs(p, exist_ok=True)
for s in ["build_gene2go.py", "fig4_qc_de_figures.R", "fig5_go_enrichment.R", "fig5_go_circle.R", "fig6_kegg.R",
          "fig7_gene_groups.R", "ppi_fetch_string.py", "fig8_ppi.R", "fig9_model.py", "fig_helpers.R", "figure_theme.R",
          "build_supplementary.py"]:
    shutil.copy2(J(ROOT, s), sub["scripts"])
shutil.copy2(J(WR, "gene_counts_integer.csv"), sub["input"])
with open(J(sub["input"], "sample_metadata.csv"), "w", newline="") as f:
    w = csv.writer(f); w.writerow(["run", "sample_title", "tissue", "genotype", "condition", "replicate"])
    for run, (tis, cond, rep, title) in sorted(key.items()):
        w.writerow([run, title, tis, "wild type", cond, rep])
for a in ["wheat_gene2go.tsv", "wheat_gene2uniprot.tsv", "wheat_gene_annotation.tsv"]:
    shutil.copy2(J(ROOT, a), sub["annotation"])
shutil.copytree(J(ROOT, "kegg_ref"), J(sub["annotation"], "kegg_ref"), dirs_exist_ok=True)
with open(UNIPROT, "rb") as fi, gzip.open(J(sub["annotation"], "uniprot_taxon4565_export.tsv.gz"), "wb") as fo:
    shutil.copyfileobj(fi, fo)
for f in glob.glob(J(ROOT, "*_DEGs*.*")) + glob.glob(J(ROOT, "*_all_genes_DESeq2.csv")) + glob.glob(J(ROOT, "*_background.txt")):
    shutil.copy2(f, sub["results"])
for fold in ["Fig4_Root_main", "Fig5_GO_Root_main", "Fig6_KEGG_Root_main", "Fig7_GeneGroups_Root_main", "Fig8_PPI_Root_main",
             "Fig9_Model_Root", "FigS_Shoot_supplementary", "FigS_GO_Shoot_supplementary", "FigS_KEGG_Shoot_supplementary",
             "FigS_GeneGroups_Shoot_supplementary", "FigS_PPI_Shoot_supplementary"]:
    dst = J(sub["results"], fold); os.makedirs(dst, exist_ok=True)
    for f in glob.glob(J(ROOT, fold, "*.csv")):
        shutil.copy2(f, dst)
shutil.copytree(J(ROOT, "ppi"), J(sub["results"], "string_networks"), dirs_exist_ok=True)
shutil.copy2(J(SRA, "salmon_QC_summary.txt"), sub["qc"])
os.makedirs(J(sub["qc"], "fastp"), exist_ok=True)
for f in glob.glob(J(SRA, "fastp_results", "*.json")):
    shutil.copy2(f, J(sub["qc"], "fastp"))
with open(J(sub["qc"], "library_qc.csv"), "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["run", "sample_title", "bioproject", "tissue", "genotype", "condition", "replicate", "raw_reads", "clean_reads",
                "retained_pct", "q30_pct", "salmon_fragments", "mapped_fragments", "mapping_rate_pct"])
    w.writerows(s3)
print("Data_and_Code assembled")
